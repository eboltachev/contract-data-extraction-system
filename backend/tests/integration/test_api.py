from pathlib import Path
from docx import Document
from openpyxl import Workbook, load_workbook
from httpx import AsyncClient, ASGITransport
import pytest
from app.main import app
from app.services.job_service import job_service

def sample_files(tmp_path: Path):
    d=Document(); d.add_paragraph('Договор № ABC-77 от 01.02.2026'); doc=tmp_path/'contract.docx'; d.save(doc)
    wb=Workbook(); ws=wb.active; ws.append(['№','Критерий','Значение']); ws.append([1,'Номер договора',None]); x=tmp_path/'template.xlsx'; wb.save(x); return doc,x

@pytest.mark.asyncio
async def test_upload_and_status_and_output(tmp_path, monkeypatch):
    monkeypatch.setattr(job_service.repo, 'root', tmp_path/'jobs')
    doc,x=sample_files(tmp_path)
    async with AsyncClient(transport=ASGITransport(app=app), base_url='http://test') as ac:
        with doc.open('rb') as df, x.open('rb') as xf:
            r=await ac.post('/api/v1/jobs', files={'contract_file':('contract.docx',df,'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),'template_file':('template.xlsx',xf,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        assert r.status_code==200
        job_id=r.json()['job_id']
        import asyncio
        for _ in range(30):
            s=await ac.get(f'/api/v1/jobs/{job_id}')
            if s.json()['status'] in ['completed','failed']: break
            await asyncio.sleep(0.1)
        assert s.json()['status']=='completed'
        dl=await ac.get(f'/api/v1/jobs/{job_id}/download')
        assert dl.status_code==200

@pytest.mark.asyncio
async def test_sse_events(tmp_path, monkeypatch):
    monkeypatch.setattr(job_service.repo, 'root', tmp_path/'jobs2')
    from app.services.progress_service import progress_service
    await progress_service.publish('x', progress=100, status='completed', agent='T', action='done')
    async with AsyncClient(transport=ASGITransport(app=app), base_url='http://test') as ac:
        r=await ac.get('/api/v1/jobs/x/events')
        assert 'data:' in r.text
